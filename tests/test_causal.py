"""The eBART causal study (analysis/causal): the frozen pre-registration, the estimators on synthetic panels with a
known effect, placebo / randomization / cluster inference, and the monthly-workbook parser.

No data/ and no network: synthetic panels, workbooks built here with openpyxl, and the committed CSVs only.
"""

import dataclasses
import hashlib
import io
import json
import math
import re
from datetime import datetime
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd
import pytest
from scipy import stats

from analysis.causal import design, did, monthly, panel, synthetic, trend_adjusted

ROOT = Path(__file__).resolve().parents[1]
CAUSAL = ROOT / "analysis" / "causal"
PREREG_SHA256 = "a05df98504818ffd4ea5fb69c720cea37eff76dd55859941acc3ce23d56e40f4"
AMENDMENT_1_SHA256 = "96395c97868f556d756415bef915f4fde13d971b5fb7b7cdb95505f876f36061"


@pytest.fixture(scope="module")
def d() -> design.Design:
    return design.load()


@pytest.fixture(scope="module")
def network_codes() -> set[str]:
    data = json.loads((ROOT / "web" / "data" / "network.json").read_text(encoding="utf-8"))
    return {s["code"] for s in data["stations"]}


# ---------------------------------------------------------------------------------------------------------------
# pre-registration


def test_prereg_file_loads_unchanged(d):
    raw = (CAUSAL / "preregistration.json").read_bytes()
    assert hashlib.sha256(raw).hexdigest() == PREREG_SHA256
    assert b"\r\n" not in raw
    assert d.sha256 == PREREG_SHA256
    assert d.fixed_on == "2026-10-09"
    assert str(d.event_date) == "2018-05-26" and str(d.week_zero) == "2018-05-28"


def test_prereg_sets_disjoint_and_complete(d, network_codes):
    design.validate(d, stations=network_codes)  # every station in exactly one set
    assert set(d.corridor) == set(d.treated) | set(d.new)
    overlap = dataclasses.replace(d, buffer=(*d.buffer, d.controls[0]))
    with pytest.raises(ValueError, match="both"):
        design.validate(overlap)
    with pytest.raises(ValueError, match="cover"):
        design.validate(d, stations=network_codes | {"XXXX"})


def test_prereg_has_34_controls(d):
    assert len(d.controls) == len(set(d.controls)) == 34
    assert set(d.controls_suburban) <= set(d.controls)
    short = dataclasses.replace(d, controls=d.controls[:-1])
    with pytest.raises(ValueError):
        design.validate(short)


def test_prereg_draws_bins_and_rules_come_from_the_file(d):
    assert d.placebo_draws["legacy"] == design.Draws(3, None)
    assert d.placebo_draws["corridor"] == design.Draws(3, None)  # Amendment 1
    assert d.placebo_draws["buffer"] == design.Draws(4, 10000)
    assert d.placebo_draws["per_station"] == design.Draws(1, None)
    assert d.seed == 20180526 and d.ci_quantiles == (0.025, 0.975) and d.caution_p == 0.10
    weekly = d.weekly_bins()
    assert weekly[:5] == [(-21, -18), (-17, -14), (-13, -10), (-9, -6), (-5, -2)]
    assert weekly[5] == (0, 3) and weekly[-1] == (88, 90) and len(weekly) == 5 + 23
    months = d.monthly_bins()
    assert months[0] == (-29, -26) and (-4, -2) in months and months[-1] == (18, 20)
    assert d.month_e(d.first_month) == -29 and d.month_e(d.last_month) == 20
    assert str(d.in_time_fake_week) == "2018-03-12" and d.week_e(d.in_time_fake_week) == -11


def test_prereg_design_py_has_no_hardcoded_station_lists(network_codes):
    for name in ("design.py", "did.py", "panel.py", "monthly.py"):
        src = (CAUSAL / name).read_text(encoding="utf-8")
        literals = set(re.findall(r"""["']([A-Z0-9]{4})["']""", src))
        assert not literals & network_codes, f"{name} hard-codes {sorted(literals & network_codes)}"


def test_amendment_1_loads_and_matches_the_original_hash(d):
    raw = (CAUSAL / "preregistration-amendment-1.json").read_bytes()
    assert hashlib.sha256(raw).hexdigest() == AMENDMENT_1_SHA256
    assert b"\r\n" not in raw
    assert json.loads(raw)["amends_sha256"] == PREREG_SHA256 == d.sha256
    assert d.amendments == (("preregistration-amendment-1.json", AMENDMENT_1_SHA256, "2026-10-10"),)
    # a missing amendment file is an error, not a silent fallback to the original
    with pytest.raises(FileNotFoundError):
        design.load.__wrapped__(design.PREREG_PATH, (CAUSAL / "no-such-amendment.json",))


def test_amendment_1_corridor_draws_are_all_sets_of_3(d):
    assert d.placebo_draws["corridor"] == design.Draws(3, None)
    assert len(did.draw_sets(34, d.placebo_draws["corridor"], d.seed)) == 5984
    assert len(did.draw_sets(16, d.placebo_draws["corridor"], d.seed)) == 560
    original = design.parse(json.loads((CAUSAL / "preregistration.json").read_bytes()))
    for k in ("legacy", "buffer", "per_station"):
        assert d.placebo_draws[k] == original.placebo_draws[k], k
    assert original.placebo_draws["corridor"] == design.Draws(5, 10000)


def test_amendment_1_rejects_a_mismatched_was():
    raw = json.loads((CAUSAL / "preregistration.json").read_bytes())
    amendment = json.loads((CAUSAL / "preregistration-amendment-1.json").read_bytes())
    bad = json.loads(json.dumps(amendment))
    bad["changes"]["inference.placebo_draws.corridor"]["was"] = "something else"
    with pytest.raises(ValueError, match="expects"):
        design.apply_amendment(raw, bad, PREREG_SHA256)
    with pytest.raises(ValueError, match="sha256"):
        design.apply_amendment(raw, amendment, "0" * 64)
    applied = design.apply_amendment(raw, amendment, PREREG_SHA256)
    assert applied["inference"]["placebo_draws"]["corridor"].startswith("all C(34,3)")
    assert raw["inference"]["placebo_draws"]["corridor"].startswith("10000")  # the original is not mutated


# ---------------------------------------------------------------------------------------------------------------
# estimators on synthetic panels


def test_synthetic_twfe_recovers_known_effect_with_low_noise(d):
    rng = np.random.default_rng(1)
    for scenario in ("moved", "new"):
        p = synthetic.generate(scenario, rng, d, sigma=0.002)
        truth = synthetic.truth(scenario, d)
        res = did.twfe(p, d.treated, d.controls, "entries")
        assert abs(res["b"] - truth["legacy"]) < 0.01
        assert res["n_treated"] == 3 and res["n_controls"] == 34 and res["n_periods"] == 111
        corridor = did.aggregate_unit(p, d.corridor, "CORRIDOR")
        assert abs(did.dd(corridor, ["CORRIDOR"], d.controls) - truth["corridor"]) < 0.01


def test_synthetic_event_study_pre_bins_zero_post_bins_effect(d):
    p = synthetic.generate("moved", np.random.default_rng(2), d, sigma=0.002)
    truth = synthetic.truth("moved", d)["legacy"]
    es = did.event_study(p, d.treated, d.controls, "entries", d.weekly_bins(), d.week_ref_bin)
    for row in es["bins"]:
        expected = truth if row["bin_start"] >= 0 else 0.0
        assert abs(row["coef"] - expected) < 0.01, row
    assert 0 <= es["pretrend_wald_p"] <= 1


def test_synthetic_scenarios_give_their_verdicts(d):
    rng = np.random.default_rng(3)
    assert did.headline(synthetic.generate("moved", rng, d, sigma=0.01), d)["verdict"] == "moved"
    assert did.headline(synthetic.generate("new", rng, d, sigma=0.01), d)["verdict"] == "new"


def test_synthetic_double_difference_equals_statsmodels(d):
    p = synthetic.generate("new", np.random.default_rng(4), d)
    res = did.twfe(p, d.treated, d.controls, "entries")
    assert did.dd(p, d.treated, d.controls, "entries") == pytest.approx(res["b"], abs=1e-8)
    bins = d.weekly_bins()
    es = did.event_study(p, d.treated, d.controls, "entries", bins, d.week_ref_bin)
    fast = did.dd_bins(p, d.treated, d.controls, "entries", bins, d.week_ref_bin)
    np.testing.assert_allclose(fast, [r["coef"] for r in es["bins"]], atol=1e-8)
    # the aggregated corridor unit too
    agg = did.aggregate_unit(p, d.corridor, "CORRIDOR")
    res = did.twfe(agg, ["CORRIDOR"], d.controls, "entries")
    assert did.dd(agg, ["CORRIDOR"], d.controls) == pytest.approx(res["b"], abs=1e-8)


def test_synthetic_aggregate_unit_sums_levels_with_missing_as_zero():
    p = pd.DataFrame(
        {
            "unit": ["A", "A", "B", "C", "C"],
            "period": ["w1", "w2", "w2", "w1", "w2"],
            "e": [-1, 0, 0, -1, 0],
            "entries": [10.0, 20.0, 5.0, 1.0, 2.0],
        }
    )
    out = did.aggregate_unit(p, ["A", "B"], "AB")
    agg = out[out["unit"] == "AB"].set_index("period")["entries"].to_dict()
    assert agg == {"w1": 10.0, "w2": 25.0}
    assert len(out) == len(p) + 2


def test_synthetic_weekly_panel_drops_transition_and_incomplete_weeks(d):
    weeks = pd.date_range(d.first_week, d.last_week, freq="7D").strftime("%Y-%m-%d")
    rows = [
        {"station_code": s, "week_start": w, "days": 7, "entries": 100, "exits": 100}
        for s in d.analysis_stations
        for w in weeks
    ]
    df = pd.DataFrame(rows)
    df.loc[(df["station_code"] == d.controls[0]) & (df["week_start"] == weeks[3]), "days"] = 6
    p = did.weekly_panel(df, d)
    kept = set(p["period"])
    assert weeks[3] not in kept and len(kept) == len(weeks) - 2  # one incomplete, one transition (e = -1)
    assert not (p["e"] == -1).any() and p["e"].min() == -21 and p["e"].max() == 90


def test_synthetic_committed_week_csv_has_every_station_and_full_new_station_weeks(d, network_codes):
    df = pd.read_csv(panel.WEEK_CSV)
    assert list(df.columns) == panel.COLUMNS
    assert set(df["station_code"]) == network_codes
    n_weeks = (d.last_week - d.first_week).days // 7 + 1
    for s in d.new + d.analysis_stations:
        assert (df["station_code"] == s).sum() == n_weeks, s
    assert (df[["days", "entries", "exits"]] >= 0).all().all()


# ---------------------------------------------------------------------------------------------------------------
# placebo / randomization / cluster inference


def test_placebo_draw_counts_and_seed_are_deterministic(d):
    n = len(d.controls)
    counts = {k: len(did.draw_sets(n, v, d.seed)) for k, v in d.placebo_draws.items()}
    assert counts == {"legacy": math.comb(34, 3), "corridor": 5984, "buffer": 10000, "per_station": 34}
    fresh = did.draw_sets.__wrapped__(n, d.placebo_draws["buffer"], d.seed)
    np.testing.assert_array_equal(fresh, did.draw_sets(n, d.placebo_draws["buffer"], d.seed))
    assert len({tuple(r) for r in fresh}) == 10000 and (np.diff(fresh, axis=1) > 0).all()
    other = did.draw_sets.__wrapped__(n, d.placebo_draws["buffer"], d.seed + 1)
    assert not np.array_equal(fresh, other)
    # more draws than combinations (e.g. 16 suburban controls): every combination instead
    assert len(did.draw_sets(16, d.placebo_draws["buffer"], d.seed)) == math.comb(16, 4) == 1820


def test_placebo_interval_and_bands_follow_the_preregistered_formulas(d):
    p = synthetic.generate("new", np.random.default_rng(5), d)
    bins = d.weekly_bins()
    r = did.placebo(
        p,
        d.treated,
        d.controls,
        "entries",
        d.placebo_draws["legacy"],
        d.seed,
        d.ci_quantiles,
        bins=bins,
        ref_bin=d.week_ref_bin,
    )
    q_lo, q_hi = np.quantile(r.draws_b, d.ci_quantiles)
    assert r.b == pytest.approx(did.dd(p, d.treated, d.controls))
    assert r.placebo_ci == pytest.approx([r.b - q_hi, r.b - q_lo])
    assert r.ri_p == pytest.approx((1 + np.sum(np.abs(r.draws_b) >= abs(r.b))) / (1 + len(r.draws_b)))
    assert r.n_draws == math.comb(34, 3)
    np.testing.assert_allclose(
        [b["coef"] for b in r.bins], did.dd_bins(p, d.treated, d.controls, "entries", bins, d.week_ref_bin)
    )
    ref = next(b for b in r.bins if (b["bin_start"], b["bin_end"]) == tuple(d.week_ref_bin))
    assert ref["coef"] == ref["lo"] == ref["hi"] == 0
    assert 0 < r.pretrend_ri_p <= 1
    # one placebo draw reproduced by hand: the first 3 controls treated, the other 31 as controls
    sets = did.draw_sets(34, d.placebo_draws["legacy"], d.seed)
    s = [d.controls[i] for i in sets[0]]
    rest = [c for c in d.controls if c not in s]
    assert r.draws_b[0] == pytest.approx(did.dd(p, s, rest), abs=1e-12)


def test_placebo_corridor_draw_matches_an_aggregated_unit(d):
    p = synthetic.generate("new", np.random.default_rng(6), d)
    r = did.placebo(
        p,
        d.corridor,
        d.controls,
        "entries",
        d.placebo_draws["corridor"],
        d.seed,
        d.ci_quantiles,
        aggregate=True,
    )
    agg = did.aggregate_unit(p, d.corridor, "CORRIDOR")
    assert r.b == pytest.approx(did.dd(agg, ["CORRIDOR"], d.controls), abs=1e-12)
    s = [d.controls[i] for i in did.draw_sets(34, d.placebo_draws["corridor"], d.seed)[7]]
    fake = did.aggregate_unit(p, s, "FAKE")
    rest = [c for c in d.controls if c not in s]
    assert r.draws_b[7] == pytest.approx(did.dd(fake, ["FAKE"], rest), abs=1e-12)


def test_randomization_p_is_minimal_under_a_large_effect(d):
    p = synthetic.generate("new", np.random.default_rng(7), d, sigma=0.01)
    r = did.placebo(
        p,
        d.corridor,
        d.controls,
        "entries",
        d.placebo_draws["corridor"],
        d.seed,
        d.ci_quantiles,
        aggregate=True,
    )
    assert r.ri_p == pytest.approx(1 / (r.n_draws + 1))
    assert r.placebo_ci[0] > 0
    p = synthetic.generate("moved", np.random.default_rng(8), d, sigma=0.01)
    r = did.placebo(p, d.treated, d.controls, "entries", d.placebo_draws["legacy"], d.seed, d.ci_quantiles)
    assert r.ri_p <= 2 / (r.n_draws + 1)


def test_randomization_null_rejection_rate_is_about_5_percent(d):
    """Legacy (3 pooled stations vs placebo sets of 3) and corridor (one aggregated unit vs summed placebo sets of 3,
    per Amendment 1: analysis/causal/preregistration-amendment-1.json)."""
    rng = np.random.default_rng(9)
    sims = 80
    rejects = {"legacy": 0, "corridor": 0}
    for _ in range(sims):
        p = synthetic.generate("null", rng, d)
        common = {"pool": d.controls, "outcome": "entries", "seed": d.seed, "quantiles": d.ci_quantiles}
        leg = did.placebo(p, d.treated, draws=d.placebo_draws["legacy"], **common)
        cor = did.placebo(p, d.corridor, draws=d.placebo_draws["corridor"], aggregate=True, **common)
        rejects["legacy"] += leg.ri_p <= 0.05
        rejects["corridor"] += cor.ri_p <= 0.05
    for k, n in rejects.items():
        assert n / sims <= 0.15, k  # loose: 80 sims, nominal 5%


def test_verdict_and_caution_rules():
    assert did.verdict([0.01, 0.2], [-0.3, -0.1]) == "new"
    assert did.verdict([-0.05, 0.05], [-0.3, -0.1]) == "moved"
    assert did.verdict([-0.2, -0.01], [-0.3, -0.1]) == "fewer"
    assert did.verdict([-0.05, 0.05], [-0.3, 0.1]) == "inconclusive"
    assert did.caution(0.09, 0.5, [-0.1, 0.1], 0.05, 0.10) is True
    assert did.caution(0.5, 0.5, [-0.1, 0.1], 0.05, 0.10) is False
    assert did.caution(0.5, 0.5, [-0.2, -0.1], 0.05, 0.10) is True  # monthly excludes 0, opposite sign
    assert did.caution(0.5, 0.5, [0.1, 0.2], 0.05, 0.10) is False
    assert did.pct(math.log(1.15)) == pytest.approx(15.0)
    imp = did.implied_daily_entries(7000.0, math.log(1.15), [0.0, math.log(1.3)])
    assert imp["point"] == pytest.approx(1000 * (1 - 1 / 1.15))
    assert imp["lo"] == 0.0 and imp["hi"] > imp["point"]


def test_cluster_cr1_matches_a_hand_computed_sandwich():
    rng = np.random.default_rng(10)
    units = [f"U{i}" for i in range(6)]
    periods = [f"p{j}" for j in range(5)]
    p = pd.DataFrame(
        [
            {
                "unit": u,
                "period": t,
                "e": j - 2,
                "entries": float(np.exp(5 + 0.3 * i + 0.1 * j + rng.normal(0, 0.1))),
            }
            for i, u in enumerate(units)
            for j, t in enumerate(periods)
        ]
    )
    treated, controls = units[:2], units[2:]
    res = did.twfe(p, treated, controls, "entries")
    # by hand: y on [D, unit dummies, period dummies (first dropped)], sorted by unit then period
    sub = p.sort_values(["unit", "period"])
    y = np.log(sub["entries"].to_numpy())
    dcol = (sub["unit"].isin(treated) & (sub["e"] >= 0)).to_numpy(float)
    ud = (sub["unit"].to_numpy()[:, None] == np.array(units)[None, :]).astype(float)
    pdm = (sub["period"].to_numpy()[:, None] == np.array(periods[1:])[None, :]).astype(float)
    x = np.column_stack([dcol, ud, pdm])
    n, k = x.shape
    bread = np.linalg.inv(x.T @ x)
    beta = bread @ x.T @ y
    u = y - x @ beta
    g = sub["unit"].to_numpy()
    meat = sum(np.outer(x[g == c].T @ u[g == c], x[g == c].T @ u[g == c]) for c in units)
    big_g = len(units)
    v = bread @ meat @ bread * big_g / (big_g - 1) * (n - 1) / (n - k)
    se = math.sqrt(v[0, 0])
    assert res["b"] == pytest.approx(beta[0], abs=1e-10)
    assert res["se"] == pytest.approx(se, rel=1e-8)
    t = stats.t.ppf(0.975, big_g - 1)
    assert res["cluster_ci"] == pytest.approx([beta[0] - t * se, beta[0] + t * se], rel=1e-8)


# ---------------------------------------------------------------------------------------------------------------
# monthly workbooks


def _workbook(sheet="Avg Weekday OD", codes=("RM", 19, "AN"), matrix=None, totals=None, date=True) -> bytes:
    """A tiny BART-style workbook: title row (with the month), header row of entry codes + "Exits", one row per exit
    station, an "Entries" totals row, and a stray summary block to the right."""
    matrix = matrix if matrix is not None else [[1.5, 20.0, 3.0], [10.0, 2.0, 30.0], [100.0, 200.0, 0.25]]
    col_sums = [sum(r[j] for r in matrix) for j in range(len(codes))]
    totals = totals if totals is not None else col_sums
    wb = openpyxl.Workbook()
    wb.active.title = "Notes"
    ws = wb.create_sheet(sheet)
    ws.append(
        [
            "Exit stations",
            "Entry stations->",
            None,
            "WEEKDAY",
            None,
            None,
            datetime(2019, 3, 1) if date else None,
        ]
    )
    ws.append([None, *codes, "Exits", None, None, "Entries"])  # stray "Entries" text in a side block
    for code, row in zip(codes, matrix, strict=True):
        ws.append([code, *row, sum(row)])
    ws.append(["Entries", *totals, sum(totals)])
    wb.create_sheet("Saturday OD")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_monthly_parser_reads_a_tiny_workbook():
    month, entries = monthly.parse_workbook(_workbook(), "Ridership_201903.xlsx")
    assert month == "2019-03"
    assert entries == pytest.approx({"RM": 111.5, "19": 222.0, "AN": 33.25})


def test_monthly_parser_older_sheet_name_and_month_from_file_name():
    month, entries = monthly.parse_workbook(
        _workbook(sheet="Weekday OD", date=False), "x/Ridership_April2016.xlsx"
    )
    assert month == "2016-04" and set(entries) == {"RM", "19", "AN"}


def test_monthly_parser_rejects_totals_that_dont_add_up():
    with pytest.raises(ValueError, match="sheet says"):
        monthly.parse_workbook(_workbook(totals=[111.5, 224.0, 33.25]), "bad.xlsx")
    monthly.parse_workbook(_workbook(totals=[111.5, 222.9, 33.25]), "rounded.xlsx")  # within 1 is fine


def test_monthly_code_map_covers_the_committed_csv_and_network(network_codes):
    codes = pd.read_csv(monthly.CODES_CSV, dtype=str)
    assert list(codes.columns) == ["code2", "station_code", "name"]
    assert codes["code2"].is_unique and codes["station_code"].is_unique
    assert set(codes["station_code"]) <= network_codes
    month_csv = pd.read_csv(monthly.MONTH_CSV)
    assert set(month_csv["station_code"]) <= set(codes["station_code"])


def test_monthly_committed_csv_is_complete(d):
    df = pd.read_csv(monthly.MONTH_CSV, dtype={"month": str})
    assert list(df.columns) == monthly.COLUMNS
    months = pd.period_range(d.first_month, d.last_month, freq="M").strftime("%Y-%m").tolist()
    assert len(months) == 50
    for s in d.analysis_stations:
        assert sorted(df.loc[df["station_code"] == s, "month"]) == months, s
    assert (df["avg_weekday_entries"] >= 0).all()
    assert not df.duplicated(["station_code", "month"]).any()


def test_monthly_panel_drops_the_transition_month(d):
    df = pd.read_csv(monthly.MONTH_CSV, dtype={"month": str})
    p = did.monthly_panel(df, d)
    assert "2018-05" not in set(p["period"]) and p["e"].min() == -29 and p["e"].max() == 20
    assert p.loc[p["period"] == "2018-06", "e"].eq(0).all()


# ---------------------------------------------------------------------------------------------------------------
# docs/FINDINGS.md against the committed findings.json

FINDINGS_DOC = ROOT / "docs" / "FINDINGS.md"
FINDINGS_JSON = ROOT / "api" / "data" / "findings.json"
FINDINGS_SECTIONS = [
    "Question",
    "Pre-registration",
    "Data",
    "Design",
    "Assumptions",
    "Results",
    "Event study",
    "Robustness and falsification",
    "Interpretation",
    "Limitations",
    "Reproduce",
]


def _pct(x: float) -> str:
    return f"{x:+.1f}%"


def test_findings_doc_has_the_sections_in_order():
    heads = re.findall(r"^## (.+)$", FINDINGS_DOC.read_text(encoding="utf-8"), flags=re.M)
    assert heads == FINDINGS_SECTIONS


def test_findings_doc_quotes_both_preregistration_hashes():
    text = FINDINGS_DOC.read_text(encoding="utf-8")
    assert PREREG_SHA256 in text and AMENDMENT_1_SHA256 in text
    doc = json.loads(FINDINGS_JSON.read_text(encoding="utf-8"))
    assert doc["preregistration"]["sha256"] == PREREG_SHA256
    assert [a["sha256"] for a in doc["preregistration"]["amendments"]] == [AMENDMENT_1_SHA256]


def test_findings_doc_images_exist():
    links = re.findall(r"!\[[^\]]*\]\(([^)]+)\)", FINDINGS_DOC.read_text(encoding="utf-8"))
    assert {Path(link).name for link in links} >= {
        "findings-event-study-weekly.png",
        "findings-event-study-monthly.png",
    }
    for link in links:
        assert (FINDINGS_DOC.parent / link).is_file(), link


def test_findings_doc_states_the_committed_headline_numbers():
    text = " ".join(FINDINGS_DOC.read_text(encoding="utf-8").replace("\u2212", "-").split())
    doc = json.loads(FINDINGS_JSON.read_text(encoding="utf-8"))
    for key in ("corridor", "legacy"):
        e = doc["estimates"][key]
        lo, hi = e["placebo_ci_pct"]
        assert _pct(e["pct"]) in text, key
        assert f"{_pct(lo)} to {_pct(hi)}" in text, key
    assert f"verdict is **{doc['verdict']}**" in text
    assert ("caution flag is off" if not doc["caution"] else "caution flag is on") in text.lower()


def test_findings_doc_limitations_cover_the_known_threats():
    text = FINDINGS_DOC.read_text(encoding="utf-8")
    section = text.split("\n## Limitations\n", 1)[1].split("\n## ", 1)[0]
    lower = " ".join(section.split()).lower()
    needles = [
        "return trip",
        "covid",
        "drift",
        "parking",
        "tri delta",
        "20 pre weeks",
        "average weekday",
        "wide side",
    ]
    for needle in needles:
        assert needle in lower, needle
    assert "spillover" in lower or "sutva" in lower


# ---------------------------------------------------------------------------------------------------------------
# post-hoc trend-adjusted estimate (analysis/causal/trend_adjusted.py), not pre-registered


def _trend_panel(d, slope: float, effect: float, seed: int = 0) -> pd.DataFrame:
    """A monthly panel: controls flat (station FE + a shared month effect), the legacy stations with a linear trend in
    event time and `effect` log points from the opening on."""
    rng = np.random.default_rng(seed)
    months = pd.period_range(d.first_month, d.last_month, freq="M").strftime("%Y-%m").tolist()
    season = rng.normal(0, 0.05, len(months))
    rows = []
    for s in list(d.controls) + list(d.treated):
        fe = rng.normal(math.log(5000), 0.5)
        for m, shared in zip(months, season, strict=True):
            e = d.month_e(m)
            y = fe + shared + rng.normal(0, 0.002)
            if s in d.treated:
                y += slope * e + effect * (e >= 0)
            rows.append({"station_code": s, "month": m, "avg_weekday_entries": math.exp(y)})
    return did.monthly_panel(pd.DataFrame(rows), d)


def _unadjusted_and_adjusted(p, d, unit) -> tuple[float, float]:
    g, _, e, _ = trend_adjusted.gaps(p, d, unit, "avg_weekday_entries")
    w = trend_adjusted.trend_weights(e.astype(float), e < 0, e >= 0)
    return float(g @ did.post_weights(e)), float(g @ w)


def test_trend_adjusted_removes_a_linear_pre_trend_with_zero_effect(d):
    p = _trend_panel(d, slope=0.003, effect=0.0)
    for unit in trend_adjusted.UNITS:
        unadjusted, adjusted = _unadjusted_and_adjusted(p, d, unit)
        assert unadjusted > 0.05, unit  # the trend alone looks like an effect to the plain double difference
        assert abs(adjusted) < 0.01, unit


def test_trend_adjusted_recovers_a_known_effect_on_top_of_a_trend(d):
    p = _trend_panel(d, slope=0.003, effect=math.log(1.1), seed=1)
    for unit in trend_adjusted.UNITS:
        _, adjusted = _unadjusted_and_adjusted(p, d, unit)
        assert abs(adjusted - math.log(1.1)) < 0.01, unit


def test_trend_adjusted_unadjusted_numbers_are_the_preregistered_ones():
    doc = json.loads(trend_adjusted.OUT.read_text(encoding="utf-8"))
    findings = json.loads(FINDINGS_JSON.read_text(encoding="utf-8"))
    for unit in trend_adjusted.UNITS:
        assert doc["weekly"][unit]["unadjusted_pct"] == pytest.approx(
            findings["estimates"][unit]["pct"], abs=1e-6
        )
        monthly_pct = findings["robustness"]["monthly"][unit]["pct"]
        assert doc["monthly"][unit]["unadjusted_pct"] == pytest.approx(monthly_pct, abs=1e-6)


def test_trend_adjusted_findings_section_quotes_the_json():
    text = FINDINGS_DOC.read_text(encoding="utf-8")
    head = "### Post-hoc: trend-adjusted estimate (not pre-registered)\n"
    assert head in text.split("\n## Limitations\n", 1)[0], "the post-hoc section comes before Limitations"
    section = text.split(head, 1)[1].split("\n## ", 1)[0]
    flat = " ".join(section.replace("\u2212", "-").replace("**", "").split())
    assert "not pre-registered" in flat and "does not change the pre-registered verdict" in flat
    doc = json.loads(trend_adjusted.OUT.read_text(encoding="utf-8"))
    assert doc["status"] == "post-hoc, not pre-registered"
    for freq in ("monthly", "weekly"):
        for unit in trend_adjusted.UNITS:
            e = doc[freq][unit]
            lo, hi = e["placebo_ci_pct"]
            assert f"{_pct(e['pct'])} ({_pct(lo)} to {_pct(hi)})" in flat, (freq, unit)
            assert _pct(e["unadjusted_pct"]) in flat, (freq, unit)
            assert f"{e['trend_per_month']:+.4f}" in flat, (freq, unit)
